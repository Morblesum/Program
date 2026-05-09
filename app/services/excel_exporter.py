"""
TOPCON眼科数据管理系统
Excel导出服务模块
"""
from typing import List, Dict, Any, Optional
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from app.core.logger import logger
from app.core.config import config


class ExcelExporter:
    """Excel导出服务"""
    
    # 默认列配置
    DEFAULT_COLUMNS = [
        {'key': 'file_created_at', 'header': '文件创建时间', 'width': 20},
        {'key': 'patient_id', 'header': '患者ID', 'width': 15},
        {'key': 'patient_no', 'header': '患者编号', 'width': 15},
        {'key': 'patient_name', 'header': '姓名', 'width': 15},
        {'key': 'exam_date', 'header': '检查时间', 'width': 20},
        {'key': 'device_model', 'header': '设备型号', 'width': 15},
        {'key': 'r_sphere', 'header': '右眼球镜(D)', 'width': 15},
        {'key': 'r_cylinder', 'header': '右眼柱镜(D)', 'width': 15},
        {'key': 'r_axis', 'header': '右眼轴位(°)', 'width': 15},
        {'key': 'r_se', 'header': '右眼等效球镜(D)', 'width': 18},
        # 右眼三次测量明细
        {'key': 'R1_sph', 'header': '右①球镜(D)', 'width': 14},
        {'key': 'R1_cyl', 'header': '右①柱镜(D)', 'width': 14},
        {'key': 'R1_axis', 'header': '右①轴位(°)', 'width': 14},
        {'key': 'R2_sph', 'header': '右②球镜(D)', 'width': 14},
        {'key': 'R2_cyl', 'header': '右②柱镜(D)', 'width': 14},
        {'key': 'R2_axis', 'header': '右②轴位(°)', 'width': 14},
        {'key': 'R3_sph', 'header': '右③球镜(D)', 'width': 14},
        {'key': 'R3_cyl', 'header': '右③柱镜(D)', 'width': 14},
        {'key': 'R3_axis', 'header': '右③轴位(°)', 'width': 14},
        {'key': 'l_sphere', 'header': '左眼球镜(D)', 'width': 15},
        {'key': 'l_cylinder', 'header': '左眼柱镜(D)', 'width': 15},
        {'key': 'l_axis', 'header': '左眼轴位(°)', 'width': 15},
        {'key': 'l_se', 'header': '左眼等效球镜(D)', 'width': 18},
        # 左眼三次测量明细
        {'key': 'L1_sph', 'header': '左①球镜(D)', 'width': 14},
        {'key': 'L1_cyl', 'header': '左①柱镜(D)', 'width': 14},
        {'key': 'L1_axis', 'header': '左①轴位(°)', 'width': 14},
        {'key': 'L2_sph', 'header': '左②球镜(D)', 'width': 14},
        {'key': 'L2_cyl', 'header': '左②柱镜(D)', 'width': 14},
        {'key': 'L2_axis', 'header': '左②轴位(°)', 'width': 14},
        {'key': 'L3_sph', 'header': '左③球镜(D)', 'width': 14},
        {'key': 'L3_cyl', 'header': '左③柱镜(D)', 'width': 14},
        {'key': 'L3_axis', 'header': '左③轴位(°)', 'width': 14},
        {'key': 'pd_distance', 'header': '远用瞳距(mm)', 'width': 16},
        {'key': 'pd_near', 'header': '近用瞳距(mm)', 'width': 16},
    ]
    
    def export_to_excel(
        self, 
        data: List[Dict[str, Any]], 
        output_path: Optional[str] = None,
        columns: Optional[List[str]] = None,
        sheet_name: Optional[str] = None
    ) -> str:
        """
        导出数据到Excel
        
        Args:
            data: 数据列表
            output_path: 输出文件路径（可选，默认使用配置）
            columns: 要导出的列名列表（可选，默认全部）
            sheet_name: 工作表名称（可选，默认使用配置）
            
        Returns:
            str: 生成的Excel文件路径
        """
        if not data:
            logger.warning("没有数据可导出")
            return None
        
        # 获取配置
        export_config = config.export
        sheet_name = sheet_name or export_config.get('sheet_name', '验光数据')
        
        # 生成默认文件名
        if not output_path:
            # 确保 exports 目录存在
            export_dir = Path('exports')
            export_dir.mkdir(parents=True, exist_ok=True)

            filename = export_config.get('default_filename', '验光数据_{date}.xlsx')
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = filename.format(date=timestamp)
            output_path = str(export_dir / filename)
        
        try:
            # 创建工作簿
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name
            
            # 确定要导出的列
            cols_to_export = self._get_columns_to_export(columns)
            
            # 设置表头
            self._set_headers(ws, cols_to_export)
            
            # 填充数据
            self._fill_data(ws, data, cols_to_export)
            
            # 应用样式
            self._apply_styles(ws, len(data), cols_to_export)
            
            # 保存文件
            wb.save(output_path)
            
            logger.info(f"✓ Excel导出成功: {output_path}")
            logger.info(f"  导出记录数: {len(data)}")
            
            return output_path
            
        except Exception as e:
            logger.error(f"✗ Excel导出失败: {e}")
            raise
    
    def _get_columns_to_export(self, columns: Optional[List[str]] = None) -> List[Dict]:
        """
        获取要导出的列配置
        
        Args:
            columns: 用户指定的列名列表
            
        Returns:
            List[Dict]: 列配置列表
        """
        if not columns:
            return self.DEFAULT_COLUMNS
        
        # 根据用户指定的列名筛选
        result = []
        for col_key in columns:
            for default_col in self.DEFAULT_COLUMNS:
                if default_col['key'] == col_key:
                    result.append(default_col)
                    break
        
        return result
    
    def _get_record_value(self, record: Dict, key: str):
        """获取记录值，支持 detail 展开列（如 R1_sph, L3_cyl）"""
        # 直接从 record 取值
        if key in record:
            return record.get(key)

        # 尝试从 details 列表中匹配
        details = record.get('details', [])
        if details:
            # 解析 key: 例如 "R1_sph" -> eye="R", list_no=1, field="sphere"
            eye = key[0]  # R 或 L
            try:
                list_no = int(key[1])  # 1, 2, 3
            except (ValueError, IndexError):
                return None

            # 映射字段名
            field_map = {
                'sph': 'sphere',
                'cyl': 'cylinder',
                'axis': 'axis',
            }
            suffix = key[3:] if len(key) > 3 else key[2:]
            field = field_map.get(suffix)

            if field:
                for d in details:
                    if d.get('eye') == eye and d.get('list_no') == list_no:
                        return d.get(field)

        return None

    def _set_headers(self, ws, columns: List[Dict]):
        """
        设置表头
        
        Args:
            ws: 工作表对象
            columns: 列配置列表
        """
        headers = [col['header'] for col in columns]
        
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
    
    def _fill_data(self, ws, data: List[Dict], columns: List[Dict]):
        """
        填充数据
        
        Args:
            ws: 工作表对象
            data: 数据列表
            columns: 列配置列表
        """
        for row_idx, record in enumerate(data, start=2):
            for col_idx, col_config in enumerate(columns, start=1):
                key = col_config['key']
                value = self._get_record_value(record, key)

                # 格式化数值
                if isinstance(value, float):
                    # 度数保留2位小数
                    if 'sphere' in key or 'cylinder' in key or 'se' in key or 'sph' in key or 'cyl' in key:
                        value = f"{value:.2f}"
                    else:
                        value = f"{value:.1f}"

                ws.cell(row=row_idx, column=col_idx, value=value)
    
    def _apply_styles(self, ws, data_rows: int, columns: List[Dict]):
        """
        应用样式

        Args:
            ws: 工作表对象
            data_rows: 数据行数
            columns: 列配置列表
        """
        col_count = len(columns)

        # 表头样式
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        # 应用表头样式
        for col_idx in range(1, col_count + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # 数据区域样式
        data_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in range(2, data_rows + 2):
            for col in range(1, col_count + 1):
                cell = ws.cell(row=row, column=col)
                cell.alignment = data_alignment

        # 设置列宽（使用实际列配置中的宽度）
        for col_idx, col_config in enumerate(columns, start=1):
            width = col_config.get('width', 15)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        
        # 冻结首行
        ws.freeze_panes = "A2"

        # 添加自动筛选
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{data_rows + 1}"
    
    def get_available_columns(self) -> List[Dict]:
        """
        获取可用的列配置
        
        Returns:
            List[Dict]: 所有可用列配置
        """
        return self.DEFAULT_COLUMNS.copy()


# 全局Excel导出器实例
excel_exporter = ExcelExporter()
